"""Resume ATS Score Checker — Streamlit + Gemini Flash."""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# "gemini-flash-latest" always points to the newest Flash model.
# Override with the GEMINI_MODEL secret/env var if you want a fixed version.
DEFAULT_MODEL = "gemini-flash-latest"
MAX_RESUME_CHARS = 15000
MAX_FILE_MB = 5

SECTION_KEYS = ["formatting", "keywords", "experience", "skills", "education", "impact"]


# ---------------------------------------------------------------- helpers
def get_secret(name: str, default: str = "") -> str:
    """Read a value from Streamlit secrets, then environment variables."""
    try:
        if name in st.secrets:
            return str(st.secrets[name])
    except Exception:
        pass  # no secrets.toml available (e.g. local run)
    return os.environ.get(name, default)


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or TXT resume."""
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = "\n".join(pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # many resumes keep skills/experience in tables
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")

    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 50:
        raise ValueError(
            "Could not read enough text from this file. If it is a scanned image PDF, "
            "please upload a text-based PDF or a DOCX instead."
        )
    return text[:MAX_RESUME_CHARS]


def build_prompt(resume_text: str, job_description: str = "") -> str:
    jd_block = (
        f"JOB DESCRIPTION (score keyword match against this):\n{job_description.strip()}"
        if job_description.strip()
        else "No job description was given. Judge against general best practice for the "
        "role the resume is clearly aiming at."
    )
    return f"""You are an expert ATS (Applicant Tracking System) and senior recruiter.
Analyse the resume below and return ONLY a JSON object, no markdown, no extra text.

Scoring rules:
- ats_score: integer 0-100, how well an ATS would parse and rank this resume.
- section_scores: integer 0-100 for each of: {", ".join(SECTION_KEYS)}.
- Be strict and honest. Do not inflate scores. Never invent facts that are not in the resume.

JSON schema:
{{
  "ats_score": 0,
  "summary": "2-3 sentence overall verdict",
  "section_scores": {{{", ".join(f'"{k}": 0' for k in SECTION_KEYS)}}},
  "strengths": ["..."],
  "missing_keywords": ["..."],
  "improvements": [
    {{"area": "short title", "issue": "what is wrong", "fix": "specific rewrite or action"}}
  ]
}}

Give 3-5 strengths, up to 10 missing keywords, and 5-8 improvements ordered by impact.

{jd_block}

RESUME:
\"\"\"
{resume_text}
\"\"\"
"""


def _clamp(value, low=0, high=100) -> int:
    try:
        return max(low, min(high, int(round(float(value)))))
    except (TypeError, ValueError):
        return low


def parse_response(raw: str) -> dict:
    """Turn the model's text into a clean, validated result dict."""
    if not raw or not raw.strip():
        raise ValueError("Empty response from the model.")
    cleaned = raw.strip()
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", cleaned, flags=re.IGNORECASE)
    start, end = cleaned.find("{"), cleaned.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("Model did not return JSON.")
    data = json.loads(cleaned[start : end + 1])
    if not isinstance(data, dict):
        raise ValueError("Model returned an unexpected format.")

    sections = data.get("section_scores") or {}
    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "area": str(item.get("area", "General")),
                    "issue": str(item.get("issue", "")),
                    "fix": str(item.get("fix", "")),
                }
            )
        elif isinstance(item, str):
            improvements.append({"area": "General", "issue": "", "fix": item})

    return {
        "ats_score": _clamp(data.get("ats_score")),
        "summary": str(data.get("summary", "")),
        "section_scores": {k: _clamp(sections.get(k)) for k in SECTION_KEYS},
        "strengths": [str(s) for s in (data.get("strengths") or [])],
        "missing_keywords": [str(k) for k in (data.get("missing_keywords") or [])],
        "improvements": improvements,
    }


def analyze_resume(api_key: str, model: str, resume_text: str, job_description: str = "") -> dict:
    """Call Gemini and return the parsed analysis. Retries once on bad JSON."""
    client = genai.Client(api_key=api_key)
    config = types.GenerateContentConfig(
        response_mime_type="application/json", temperature=0.2
    )
    prompt = build_prompt(resume_text, job_description)
    last_error = None
    for _ in range(2):
        response = client.models.generate_content(model=model, contents=prompt, config=config)
        try:
            return parse_response(response.text)
        except (ValueError, json.JSONDecodeError) as err:
            last_error = err
    raise ValueError(f"Could not read the AI response: {last_error}")


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Weak"


def build_report(result: dict) -> str:
    lines = [f"ATS SCORE: {result['ats_score']}/100 ({score_label(result['ats_score'])})", ""]
    lines += [result["summary"], "", "SECTION SCORES"]
    lines += [f"- {k.title()}: {v}/100" for k, v in result["section_scores"].items()]
    lines += ["", "STRENGTHS"] + [f"- {s}" for s in result["strengths"]]
    lines += ["", "MISSING KEYWORDS", ", ".join(result["missing_keywords"]) or "None", "", "IMPROVEMENTS"]
    for i, imp in enumerate(result["improvements"], 1):
        lines += [f"{i}. {imp['area']}", f"   Issue: {imp['issue']}", f"   Fix: {imp['fix']}"]
    return "\n".join(lines)


# --------------------------------------------------------------------- UI
def main():
    st.set_page_config(page_title="Resume ATS Checker", page_icon="📄", layout="centered")
    st.title("📄 Resume ATS Score Checker")
    st.caption("Upload your resume to get an ATS score and clear improvements. Powered by Gemini Flash.")

    api_key = get_secret("GEMINI_API_KEY")
    model = get_secret("GEMINI_MODEL", DEFAULT_MODEL)
    if not api_key:
        api_key = st.text_input("Gemini API key", type="password", help="Get a free key at aistudio.google.com")

    uploaded = st.file_uploader("Upload resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, improves keyword matching)", height=150
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key.")
            return
        if uploaded.size > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
            return
        try:
            text = extract_text(uploaded.name, uploaded.getvalue())
        except Exception as err:
            st.error(f"Could not read the file: {err}")
            return
        try:
            with st.spinner("Analyzing your resume..."):
                st.session_state["result"] = analyze_resume(api_key, model, text, job_description)
        except Exception as err:
            st.error(f"Analysis failed: {err}")
            return

    result = st.session_state.get("result")
    if not result:
        return

    st.divider()
    score = result["ats_score"]
    col1, col2 = st.columns([1, 2])
    col1.metric("ATS Score", f"{score}/100", score_label(score))
    col2.write(result["summary"])
    st.progress(score / 100)

    st.subheader("Section scores")
    cols = st.columns(3)
    for i, (key, val) in enumerate(result["section_scores"].items()):
        cols[i % 3].metric(key.title(), f"{val}/100")

    if result["strengths"]:
        st.subheader("✅ Strengths")
        for s in result["strengths"]:
            st.write(f"- {s}")

    if result["missing_keywords"]:
        st.subheader("🔑 Missing keywords")
        st.write(", ".join(result["missing_keywords"]))

    st.subheader("🛠️ Improvements")
    for i, imp in enumerate(result["improvements"], 1):
        with st.expander(f"{i}. {imp['area']}", expanded=i <= 3):
            if imp["issue"]:
                st.markdown(f"**Issue:** {imp['issue']}")
            st.markdown(f"**Fix:** {imp['fix']}")

    st.download_button("Download report", build_report(result), file_name="ats_report.txt")


if __name__ == "__main__":
    main()
